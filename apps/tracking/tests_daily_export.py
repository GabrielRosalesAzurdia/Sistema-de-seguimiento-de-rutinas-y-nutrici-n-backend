import csv
import io
from datetime import date, datetime, time, timezone as dt_timezone
from unittest import mock

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.members.models import Member, User
from apps.routines.models import Routine, RoutineCategory
from .daily_export import (
    build_daily_workbook, build_key_rows, build_participant_codes, resolve_daily_range,
)
from .models import DailyNutritionLog, NutritionCheckStatus, WorkoutSessionLog
from .services import InvalidStudyRange, compute_study_metrics

OCT = lambda d: date(2026, 10, d)  # noqa: E731
RANGE = (OCT(1), OCT(31))


def local_dt(day, hour=12, minute=0):
    return timezone.make_aware(datetime.combine(day, time(hour, minute)))


class DailyExportBase(TestCase):
    """Fija "hoy" en el 31 de octubre de 2026 para que el cutoff no dependa del reloj."""

    def setUp(self):
        patcher = mock.patch("django.utils.timezone.localdate", return_value=OCT(31))
        self.today = patcher.start()
        self.addCleanup(patcher.stop)
        self.routine = Routine.objects.create(category=RoutineCategory.PECHO, estimated_calories=300)
        self.cardio = Routine.objects.create(category=RoutineCategory.CARDIO, estimated_calories=200)
        self._n = 0

    def make_member(self, first, last, start=OCT(1), planned=12, is_active=True):
        self._n += 1
        user = User.objects.create_user(
            username=f"m{self._n}@test.com", email=f"m{self._n}@test.com", password="pass1234"
        )
        member = Member.objects.create(
            user=user, first_name=first, first_last_name=last, age=25, height_cm="170.0",
            planned_training_days=planned, planned_nutrition_days=planned,
            participates_in_study=True, start_date=start, is_active=is_active,
        )
        Member.objects.filter(pk=member.pk).update(created_at=local_dt(start, 0, 1))
        member.refresh_from_db()
        return member

    def log_workout(self, member, day, minutes, hour=12, minute=0, routine=None):
        log = WorkoutSessionLog.objects.create(
            member=member, routine=routine or self.routine, duration_minutes=minutes,
        )
        WorkoutSessionLog.objects.filter(pk=log.pk).update(completed_at=local_dt(day, hour, minute))

    def log_nutrition(self, member, day, status):
        DailyNutritionLog.objects.create(member=member, date=day, status=status)

    def workbook(self, start=RANGE[0], end=RANGE[1]):
        return load_workbook(io.BytesIO(build_daily_workbook(start, end)))

    @staticmethod
    def rows_by_code(ws):
        header = [c.value for c in ws[1]]
        return header, {r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=2)}


class SheetStructureTests(DailyExportBase):
    def test_sheet_order_and_headers(self):
        self.make_member("Qwerty", "Uno")
        wb = self.workbook()
        self.assertEqual(
            wb.sheetnames,
            ["Entrenamiento 1-15", "Entrenamiento 16-31", "Nutricion 1-15", "Nutricion 16-31", "Resumen"],
        )
        header, _ = self.rows_by_code(wb["Entrenamiento 1-15"])
        self.assertEqual(header[0], "Participante")
        self.assertEqual(header[1:16], list(range(1, 16)))
        self.assertEqual(header[16:], ["Días entrenados", "Minutos totales"])
        header, _ = self.rows_by_code(wb["Nutricion 16-31"])
        self.assertEqual(header[1:17], list(range(16, 32)))
        self.assertEqual(
            header[17:], ["Días con registro", "C (HECHO)", "P (PARCIALMENTE)", "N (SE_ME_FUE)"]
        )

    def test_second_half_follows_month_length(self):
        self.make_member("Qwerty", "Uno", start=date(2026, 11, 1))
        wb = self.workbook(date(2026, 11, 1), date(2026, 11, 30))
        self.assertIn("Entrenamiento 16-30", wb.sheetnames)
        self.assertIn("Nutricion 16-30", wb.sheetnames)

    def test_default_range_is_october_2026(self):
        self.assertEqual(resolve_daily_range(None, None), (OCT(1), OCT(31)))
        self.assertEqual(resolve_daily_range("2026-10-05", None), (OCT(5), OCT(31)))

    def test_range_crossing_months_is_rejected(self):
        with self.assertRaises(InvalidStudyRange):
            resolve_daily_range("2026-10-20", "2026-11-05")

    def test_inverted_range_is_rejected(self):
        with self.assertRaises(InvalidStudyRange):
            resolve_daily_range("2026-10-20", "2026-10-05")


class TotalsReconcileTests(DailyExportBase):
    """Los totales de las hojas diarias cuadran con Resumen y con compute_study_metrics."""

    def setUp(self):
        super().setUp()
        self.member = self.make_member("Qwerty", "Uno")
        # Dos sesiones el 3 (una de cardio), una el 5 y una el 20.
        self.log_workout(self.member, OCT(3), 40)
        self.log_workout(self.member, OCT(3), 20, hour=18, routine=self.cardio)
        self.log_workout(self.member, OCT(5), 30)
        self.log_workout(self.member, OCT(20), 45)
        self.log_nutrition(self.member, OCT(1), NutritionCheckStatus.HECHO)
        self.log_nutrition(self.member, OCT(2), NutritionCheckStatus.PARCIALMENTE)
        self.log_nutrition(self.member, OCT(3), NutritionCheckStatus.HECHO)
        self.log_nutrition(self.member, OCT(20), NutritionCheckStatus.SE_ME_FUE)

    def test_training_totals_match_summary_and_study_metrics(self):
        wb = self.workbook()
        row = next(m for m in compute_study_metrics(*RANGE) if m["member"] == self.member)

        _, h1 = self.rows_by_code(wb["Entrenamiento 1-15"])
        _, h2 = self.rows_by_code(wb["Entrenamiento 16-31"])
        days = h1["P01"][-2] + h2["P01"][-2]
        minutes = h1["P01"][-1] + h2["P01"][-1]
        self.assertEqual(days, 3)
        self.assertEqual(minutes, 135)

        header, summary = self.rows_by_code(wb["Resumen"])
        s = dict(zip(header, summary["P01"]))
        self.assertEqual(days, s["Sesiones completadas"])
        self.assertEqual(days, row["completed"])
        self.assertEqual(s["Nº sesiones (control)"], 4)
        self.assertEqual(s["Minutos totales (control)"], minutes)
        # Duración promedio actual (por sesión) = minutos totales / nº de sesiones.
        self.assertEqual(round(minutes / s["Nº sesiones (control)"], 1), s["VD1 - Duración promedio (min)"])
        self.assertEqual(s["VD1 - Duración promedio (min)"], row["vd1_avg_duration"])

    def test_nutrition_totals_match_summary_and_vd2(self):
        wb = self.workbook()
        row = next(m for m in compute_study_metrics(*RANGE) if m["member"] == self.member)

        _, h1 = self.rows_by_code(wb["Nutricion 1-15"])
        _, h2 = self.rows_by_code(wb["Nutricion 16-31"])
        total = h1["P01"][-4] + h2["P01"][-4]
        c = h1["P01"][-3] + h2["P01"][-3]
        p = h1["P01"][-2] + h2["P01"][-2]
        n = h1["P01"][-1] + h2["P01"][-1]
        self.assertEqual((total, c, p, n), (4, 2, 1, 1))
        self.assertEqual(c + p + n, total)

        header, summary = self.rows_by_code(wb["Resumen"])
        s = dict(zip(header, summary["P01"]))
        self.assertEqual(total, s["Días con registro nutricional"])
        self.assertEqual(total, row["days_with_log"])
        self.assertEqual(s["Días activos"], 31)
        self.assertEqual(s["VD2 %"], round(total / s["Días activos"] * 100, 1))
        self.assertEqual(s["VD2 %"], row["vd2"])

    def test_summary_values_equal_csv_export(self):
        coach = User.objects.create_user(username="c@test.com", email="c@test.com", password="x", is_staff=True)
        client = APIClient()
        client.force_authenticate(coach)
        csv_rows = list(csv.reader(io.StringIO(
            client.get("/api/tracking/study-export/?start=2026-10-01&end=2026-10-31").content.decode()
        )))
        wb = self.workbook()
        header, summary = self.rows_by_code(wb["Resumen"])
        xlsx_values = summary["P01"][:13]
        csv_values = csv_rows[1][1:]
        self.assertEqual(header[1:13], csv_rows[0][1:])
        for xv, cv in zip(xlsx_values[1:], csv_values):
            if cv == "":
                self.assertIsNone(xv)
            else:
                self.assertEqual(float(xv), float(cv))

    def test_cell_contents(self):
        wb = self.workbook()
        header, h1 = self.rows_by_code(wb["Entrenamiento 1-15"])
        row = h1["P01"]
        self.assertEqual(row[3], 60)  # día 3: 40 + 20 minutos
        self.assertEqual(row[5], 30)
        self.assertIsNone(row[4])  # día 4 sin entrenar: vacía, no 0
        _, n1 = self.rows_by_code(wb["Nutricion 1-15"])
        self.assertEqual((n1["P01"][1], n1["P01"][2], n1["P01"][4]), ("C", "P", None))
        _, n2 = self.rows_by_code(wb["Nutricion 16-31"])
        self.assertEqual(n2["P01"][5], "N")  # día 20


class SameDayAndTimezoneTests(DailyExportBase):
    def test_two_sessions_same_day_count_once_and_sum_minutes(self):
        member = self.make_member("Qwerty", "Uno")
        self.log_workout(member, OCT(7), 25)
        self.log_workout(member, OCT(7), 35, hour=19, routine=self.cardio)
        wb = self.workbook()
        _, h1 = self.rows_by_code(wb["Entrenamiento 1-15"])
        self.assertEqual(h1["P01"][7], 60)
        self.assertEqual(h1["P01"][-2], 1)  # un solo día entrenado
        self.assertEqual(h1["P01"][-1], 60)

    def test_late_night_session_stays_on_local_day(self):
        member = self.make_member("Qwerty", "Uno")
        # 23:30 en Guatemala del 15 = 05:30 UTC del 16: debe quedar en el día 15.
        self.log_workout(member, OCT(15), 50, hour=23, minute=30)
        # 00:10 en Guatemala del 16 = 06:10 UTC del 16: día 16.
        self.log_workout(member, OCT(16), 20, hour=0, minute=10)
        stored = WorkoutSessionLog.objects.order_by("completed_at").first().completed_at
        self.assertEqual(stored.astimezone(dt_timezone.utc).date(), OCT(16))  # en UTC ya es el 16

        wb = self.workbook()
        _, h1 = self.rows_by_code(wb["Entrenamiento 1-15"])
        _, h2 = self.rows_by_code(wb["Entrenamiento 16-31"])
        self.assertEqual(h1["P01"][15], 50)
        self.assertEqual(h1["P01"][-2], 1)
        self.assertEqual(h2["P01"][1], 20)  # día 16 = primera columna de la segunda hoja
        self.assertEqual(h2["P01"][-2], 1)


class MidMonthJoinTests(DailyExportBase):
    def test_days_before_start_and_after_cutoff_are_dashes(self):
        self.today.return_value = OCT(20)
        member = self.make_member("Qwerty", "Uno", start=OCT(10))
        self.log_workout(member, OCT(12), 30)
        self.log_nutrition(member, OCT(12), NutritionCheckStatus.HECHO)
        wb = self.workbook()
        _, h1 = self.rows_by_code(wb["Entrenamiento 1-15"])
        _, h2 = self.rows_by_code(wb["Entrenamiento 16-31"])
        self.assertEqual(h1["P01"][1:10], ["-"] * 9)  # días 1-9: no estaba activo
        self.assertIsNone(h1["P01"][10])  # día 10: activo, sin sesión
        self.assertEqual(h1["P01"][12], 30)
        self.assertEqual(h2["P01"][1:6], [None] * 5)  # días 16-20: activo, sin sesión
        self.assertEqual(h2["P01"][6:17], ["-"] * 11)  # días 21-31: futuro
        _, n1 = self.rows_by_code(wb["Nutricion 1-15"])
        self.assertEqual(n1["P01"][1:10], ["-"] * 9)

    def test_variation_blank_with_less_than_two_weeks_window(self):
        member = self.make_member("Qwerty", "Uno", start=OCT(27))  # ventana de 5 días
        self.log_workout(member, OCT(28), 30)
        self.log_nutrition(member, OCT(28), NutritionCheckStatus.HECHO)
        wb = self.workbook()
        header, summary = self.rows_by_code(wb["Resumen"])
        s = dict(zip(header, summary["P01"]))
        self.assertEqual(s["Días activos"], 5)
        self.assertIsNone(s["VD1 - Variación %"])
        self.assertIsNone(s["VD2 - Variación %"])

    def test_variation_present_with_two_week_window(self):
        member = self.make_member("Qwerty", "Uno", start=OCT(1))
        self.log_workout(member, OCT(2), 30)
        wb = self.workbook()
        header, summary = self.rows_by_code(wb["Resumen"])
        s = dict(zip(header, summary["P01"]))
        self.assertIsNotNone(s["VD1 - Variación %"])


class CodesInactiveAndPermissionTests(DailyExportBase):
    def test_codes_follow_start_date_then_id_and_are_stable(self):
        late = self.make_member("Late", "Uno", start=OCT(10))
        early = self.make_member("Early", "Dos", start=OCT(2))
        same_a = self.make_member("SameA", "Tres", start=OCT(5))
        same_b = self.make_member("SameB", "Cuatro", start=OCT(5))
        order = [m.id for _, m in build_participant_codes()]
        self.assertEqual(order, [early.id, same_a.id, same_b.id, late.id])
        self.assertEqual(order, [m.id for _, m in build_participant_codes()])

    def test_inactive_excluded_from_sheets_but_kept_in_key_without_shifting_codes(self):
        first = self.make_member("Primero", "Uno", start=OCT(1))
        gone = self.make_member("Baja", "Dos", start=OCT(2), is_active=False)
        third = self.make_member("Tercero", "Tres", start=OCT(3))
        self.log_workout(gone, OCT(5), 30)

        wb = self.workbook()
        for name in ("Entrenamiento 1-15", "Nutricion 1-15", "Resumen"):
            _, rows = self.rows_by_code(wb[name])
            self.assertEqual(list(rows), ["P01", "P03"])  # P02 (inactivo) no aparece ni se reasigna

        key = {r[0]: r for r in build_key_rows()}
        self.assertEqual(key["P02"][3], "inactivo")
        self.assertEqual(key["P03"][1], third.full_name)
        self.assertEqual(key["P01"][3], "activo")
        self.assertEqual(first.full_name, key["P01"][1])

    def test_xlsx_contains_no_member_names(self):
        member = self.make_member("Qwerty", "Zxcvbn")
        self.log_workout(member, OCT(3), 30)
        wb = self.workbook()
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for value in row:
                    self.assertNotIn("Qwerty", str(value))
                    self.assertNotIn("Zxcvbn", str(value))

    def test_endpoints_require_coach(self):
        member_user = User.objects.create_user(username="u@test.com", email="u@test.com", password="x")
        coach = User.objects.create_user(username="c@test.com", email="c@test.com", password="x", is_staff=True)
        self.make_member("Qwerty", "Uno")

        anon = APIClient()
        client = APIClient()
        client.force_authenticate(member_user)
        for url in ("/api/tracking/study-export-daily/", "/api/tracking/study-export-key/"):
            self.assertIn(anon.get(url).status_code, (401, 403))
            self.assertEqual(client.get(url).status_code, 403)

        client.force_authenticate(coach)
        resp = client.get("/api/tracking/study-export-daily/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("spreadsheetml", resp["Content-Type"])
        self.assertIn("estudio_diario.xlsx", resp["Content-Disposition"])
        resp = client.get("/api/tracking/study-export-key/")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("clave_participantes.csv", resp["Content-Disposition"])
        self.assertIn("Qwerty", resp.content.decode())

    def test_cross_month_range_returns_400(self):
        coach = User.objects.create_user(username="c@test.com", email="c@test.com", password="x", is_staff=True)
        client = APIClient()
        client.force_authenticate(coach)
        resp = client.get("/api/tracking/study-export-daily/?start=2026-10-20&end=2026-11-05")
        self.assertEqual(resp.status_code, 400)
        self.assertIn("mismo mes", resp.content.decode())
